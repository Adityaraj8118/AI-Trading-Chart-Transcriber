from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

HEADERS = [
    "Slide_No",
    "Ticker",
    "Timeframe",
    "Price",
    "Trend",
    "Pattern",
    "RSI_14D",
    "Relation_to_20_SMA",
    "Volume_Analysis",
    "Immediate_Support",
    "Immediate_Resistance",
    "RSI",
    "BB_Status",
    "Divergence",
    "MACD",
    "Moving_Averages",
    "Candlestick_Pattern",
    "Prediction",
    "Confidence",
    "Remarks"
]


def save_excel(data_list, output_file):

    wb = Workbook()

    ws = wb.active

    ws.title = "Trading Analysis"

    # Header
    for col, header in enumerate(HEADERS, start=1):

        cell = ws.cell(row=1, column=col)

        cell.value = header

        cell.font = Font(bold=True)

        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="CCE5FF"
        )

    # Data
    row = 2

    for item in data_list:

        for col, header in enumerate(HEADERS, start=1):

            ws.cell(
                row=row,
                column=col
            ).value = item.get(header, "")

        row += 1

    # Auto Width
    for column_cells in ws.columns:

        length = max(len(str(cell.value or "")) for cell in column_cells)

        ws.column_dimensions[
            get_column_letter(column_cells[0].column)
        ].width = length + 5

    ws.freeze_panes = "A2"

    ws.auto_filter.ref = ws.dimensions

    wb.save(output_file)

    print(f"\nExcel Saved : {output_file}")